"""Génère docs/Modele-Agenda-GestionAgenda.xlsx, aligné sur FeuilleAgenda.cs (colonnes, listes, libellés)."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

LIGNES = 2000
SORTIE = Path(__file__).with_name("Modele-Agenda-GestionAgenda.xlsx")

ENTETES = [
    "ID application", "Date", "Jour", "Heure début", "Heure fin", "Durée", "Type", "Objet", "Société",
    "Dirigeant", "Lieu ou lien visio", "Participants", "Demandeur", "Priorité", "Statut",
    "Documents à préparer", "Brief prêt ?", "Logistique", "Notes et décisions",
]
TYPES = ["Rendez-vous", "Réunion", "Appel", "Visio", "Déjeuner", "Déplacement", "Événement", "Temps bloqué", "Personnel"]
PRIORITES = ["Basse", "Moyenne", "Haute", "Critique"]
STATUTS = ["À confirmer", "Accepté", "Confirmé", "Reporté", "Refusé", "Annulé", "Terminé"]
OUI_NON = ["Oui", "Non"]
LARGEURS = [14, 12, 11, 11, 11, 9, 14, 38, 20, 22, 26, 40, 22, 11, 13, 28, 11, 24, 40]

APPLICATION = PatternFill("solid", fgColor="1F3A5F")
FORMULE = PatternFill("solid", fgColor="5B6B82")
ASSISTANTE = PatternFill("solid", fgColor="2E7D5B")
BLANC = Font(bold=True, color="FFFFFF")
GRIS_ID = PatternFill("solid", fgColor="F1F3F6")
BORDURE = Border(bottom=Side(style="thin", color="D0D5DD"))

classeur = Workbook()

# Onglet Listes : source des listes déroulantes.
listes = classeur.active
listes.title = "Listes"
for colonne, (titre, valeurs) in enumerate(
    [("Type", TYPES), ("Priorité", PRIORITES), ("Statut", STATUTS), ("Oui / Non", OUI_NON)], start=1
):
    lettre = get_column_letter(colonne)
    listes[f"{lettre}1"] = titre
    listes[f"{lettre}1"].font = BLANC
    listes[f"{lettre}1"].fill = APPLICATION
    listes.column_dimensions[lettre].width = 16
    for ligne, valeur in enumerate(valeurs, start=2):
        listes[f"{lettre}{ligne}"] = valeur

# Onglet Agenda : doit porter le nom configuré dans GoogleSheets:Onglet.
agenda = classeur.create_sheet("Agenda", 0)
for colonne, (titre, largeur) in enumerate(zip(ENTETES, LARGEURS), start=1):
    cellule = agenda.cell(row=1, column=colonne, value=titre)
    cellule.font = BLANC
    cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cellule.fill = FORMULE if colonne in (3, 6) else APPLICATION if colonne <= 15 else ASSISTANTE
    agenda.column_dimensions[get_column_letter(colonne)].width = largeur
agenda.row_dimensions[1].height = 32
agenda.freeze_panes = "C2"
agenda.auto_filter.ref = f"A1:S{LIGNES + 1}"

agenda["A1"].comment = Comment(
    "Rempli par l'application. Ne pas modifier : c'est ce qui relie la ligne à la demande.", "GestionAgenda"
)
agenda["O1"].comment = Comment(
    "Colonnes A à O : écrites par l'application à chaque changement. "
    "Colonnes P à S : libres, jamais écrasées.", "GestionAgenda"
)

for ligne in range(2, LIGNES + 2):
    agenda[f"A{ligne}"].fill = GRIS_ID
    agenda[f"A{ligne}"].font = Font(color="8A94A6", size=8)
    agenda[f"B{ligne}"].number_format = "dd/mm/yyyy"
    agenda[f"C{ligne}"] = (
        f'=IF(B{ligne}="","",CHOOSE(WEEKDAY(B{ligne},2),'
        '"Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi","Dimanche"))'
    )
    agenda[f"D{ligne}"].number_format = "hh:mm"
    agenda[f"E{ligne}"].number_format = "hh:mm"
    agenda[f"F{ligne}"] = f'=IF(OR(D{ligne}="",E{ligne}=""),"",E{ligne}-D{ligne})'
    agenda[f"F{ligne}"].number_format = "[h]:mm"
    for lettre in "HLPRS":
        agenda[f"{lettre}{ligne}"].alignment = Alignment(wrap_text=True, vertical="top")
    for colonne in range(1, 20):
        agenda.cell(row=ligne, column=colonne).border = BORDURE


def liste(plage_source: str, cible: str, message: str) -> None:
    validation = DataValidation(type="list", formula1=f"=Listes!{plage_source}", allow_blank=True)
    validation.error = message
    validation.errorTitle = "Valeur non reconnue"
    validation.showErrorMessage = True
    agenda.add_data_validation(validation)
    validation.add(cible)


fin = LIGNES + 1
liste(f"$A$2:$A${len(TYPES) + 1}", f"G2:G{fin}", "Choisissez un type de la liste : l'application ne reconnaît que ceux-là.")
liste(f"$B$2:$B${len(PRIORITES) + 1}", f"N2:N{fin}", "Choisissez une priorité de la liste.")
liste(f"$C$2:$C${len(STATUTS) + 1}", f"O2:O{fin}", "Choisissez un statut de la liste.")
liste(f"$D$2:$D${len(OUI_NON) + 1}", f"Q2:Q{fin}", "Oui ou Non.")

date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
date.error = "Saisissez une date (jj/mm/aaaa)."
date.showErrorMessage = True
agenda.add_data_validation(date)
date.add(f"B2:B{fin}")

heure = DataValidation(type="time", operator="between", formula1="0", formula2="0.999988", allow_blank=True)
heure.error = "Saisissez une heure (hh:mm)."
heure.showErrorMessage = True
agenda.add_data_validation(heure)
heure.add(f"D2:E{fin}")

zone = f"A2:S{fin}"
regles = [
    ('$O2="Annulé"', PatternFill("solid", fgColor="E5E7EB"), Font(color="6B7280", strike=True)),
    ('$O2="Refusé"', PatternFill("solid", fgColor="E5E7EB"), Font(color="6B7280")),
    ('$O2="Terminé"', PatternFill("solid", fgColor="F3F4F6"), Font(color="4B5563")),
    ('$O2="À confirmer"', PatternFill("solid", fgColor="FEF3C7"), None),
    ('$O2="Reporté"', PatternFill("solid", fgColor="FFEDD5"), None),
    ('$O2="Accepté"', PatternFill("solid", fgColor="DBEAFE"), None),
    ('OR($G2="Déplacement",$G2="Temps bloqué",$G2="Personnel",$G2="Événement")',
     PatternFill("solid", fgColor="EEF0F3"), Font(color="374151", italic=True)),
]
for formule, remplissage, police in regles:
    agenda.conditional_formatting.add(
        zone, FormulaRule(formula=[formule], fill=remplissage, font=police, stopIfTrue=True)
    )
agenda.conditional_formatting.add(
    f"N2:N{fin}",
    FormulaRule(formula=['OR($N2="Haute",$N2="Critique")'], font=Font(bold=True, color="B42318")),
)
agenda.conditional_formatting.add(
    f"Q2:Q{fin}",
    FormulaRule(formula=['AND($Q2="Non",$B2<>"",$B2-TODAY()<=1,$B2>=TODAY())'],
                fill=PatternFill("solid", fgColor="FEE4E2"), font=Font(bold=True, color="B42318")),
)

# Onglet Suivi : actions décidées en réunion, entièrement libre.
suivi = classeur.create_sheet("Suivi", 1)
for colonne, (titre, largeur) in enumerate(
    zip(["Date de la réunion", "Objet", "Action", "Responsable", "Échéance", "Fait ?"], [16, 34, 46, 22, 14, 10]),
    start=1,
):
    cellule = suivi.cell(row=1, column=colonne, value=titre)
    cellule.font = BLANC
    cellule.fill = ASSISTANTE
    cellule.alignment = Alignment(horizontal="center", vertical="center")
    suivi.column_dimensions[get_column_letter(colonne)].width = largeur
suivi.freeze_panes = "A2"
for ligne in range(2, 502):
    suivi[f"A{ligne}"].number_format = "dd/mm/yyyy"
    suivi[f"E{ligne}"].number_format = "dd/mm/yyyy"
    suivi[f"C{ligne}"].alignment = Alignment(wrap_text=True, vertical="top")
fait = DataValidation(type="list", formula1=f"=Listes!$D$2:$D${len(OUI_NON) + 1}", allow_blank=True)
suivi.add_data_validation(fait)
fait.add("F2:F501")
suivi.conditional_formatting.add(
    "A2:F501",
    FormulaRule(formula=['AND($E2<>"",$E2<TODAY(),$F2<>"Oui")'], font=Font(bold=True, color="B42318")),
)
suivi.conditional_formatting.add(
    "A2:F501", FormulaRule(formula=['$F2="Oui"'], font=Font(color="6B7280", strike=True))
)

# Onglet Légende.
legende = classeur.create_sheet("Légende")
legende.column_dimensions["A"].width = 26
legende.column_dimensions["B"].width = 90
lignes_legende = [
    ("En-tête bleu foncé", "Colonnes A, B, D, E, G à O : écrites par l'application. Toute modification manuelle sera écrasée.", APPLICATION),
    ("En-tête gris", "Colonnes C et F : calculées automatiquement (jour, durée).", FORMULE),
    ("En-tête vert", "Colonnes P à S et onglet Suivi : libres pour l'assistante, jamais modifiées par l'application.", ASSISTANTE),
    ("Import", "Une ligne sans ID, datée dans le futur, peut être importée depuis l'Agenda de l'application "
               "(bouton « Importer les lignes à venir »). L'ID est alors inscrit en colonne A.", None),
    ("Types", "Déplacement, Temps bloqué, Personnel et Événement deviennent des indisponibilités du dirigeant ; "
              "les autres types deviennent des demandes de rencontre.", None),
    ("Participants", "Séparés par « ; ». Format : Nom <email>, Nom seul (membre de la société) ou email seul.", None),
    ("Société / Dirigeant", "Doivent correspondre aux noms saisis dans l'application.", None),
]
for ligne, (titre, texte, couleur) in enumerate(lignes_legende, start=1):
    legende[f"A{ligne}"] = titre
    legende[f"B{ligne}"] = texte
    legende[f"B{ligne}"].alignment = Alignment(wrap_text=True, vertical="top")
    legende[f"A{ligne}"].font = BLANC if couleur else Font(bold=True)
    if couleur:
        legende[f"A{ligne}"].fill = couleur

classeur.save(SORTIE)
print(f"Modèle écrit : {SORTIE}")
